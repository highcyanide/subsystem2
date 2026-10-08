import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_mint_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # --- Mint-Green & Clean White Palette ---
    BG_WHITE = RGBColor(248, 250, 252)       # Soft off-white / light slate 50
    CARD_WHITE = RGBColor(255, 255, 255)     # Pure white for cards
    MINT_LIGHT = RGBColor(236, 253, 245)     # Mint 50
    MINT_ACCENT = RGBColor(16, 185, 129)     # Mint / Emerald 500
    MINT_DARK = RGBColor(5, 150, 105)        # Mint / Emerald 600
    MINT_BORDER = RGBColor(167, 243, 208)    # Mint 200 border
    TEXT_DARK = RGBColor(15, 23, 42)         # Deep slate 900 for high readability
    TEXT_MUTED = RGBColor(71, 85, 105)       # Slate 600 for body descriptions
    TEXT_SUBTLE = RGBColor(100, 116, 139)    # Slate 500
    TEAL = RGBColor(13, 148, 136)            # Teal 600
    AMBER_BG = RGBColor(254, 243, 199)       # Amber 100
    AMBER_TEXT = RGBColor(180, 83, 9)        # Amber 700
    ROSE_BG = RGBColor(254, 226, 226)        # Rose 100
    ROSE_TEXT = RGBColor(185, 28, 28)        # Rose 700

    def set_white_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_WHITE
        bg.line.fill.background()
        
        # Soft top mint branding banner line
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.12))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = MINT_ACCENT
        top_bar.line.fill.background()
        return bg

    def add_clean_header(slide, title_text, section_tag="SUBSYSTEM 1: WAREHOUSE INVENTORY"):
        # Section pill
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(4.5), Inches(0.32))
        pill.fill.solid()
        pill.fill.fore_color.rgb = MINT_LIGHT
        pill.line.color.rgb = MINT_BORDER
        pill.line.width = Pt(1)
        tf_pill = pill.text_frame
        tf_pill.word_wrap = True
        p_p = tf_pill.paragraphs[0]
        p_p.text = section_tag.upper()
        p_p.font.size = Pt(10)
        p_p.font.bold = True
        p_p.font.color.rgb = MINT_DARK
        p_p.alignment = PP_ALIGN.CENTER

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_DARK

        # Soft divider
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.42), Inches(11.733), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = MINT_BORDER
        line.line.fill.background()

    def add_white_card(slide, left, top, width, height, title="", title_color=MINT_DARK, bg_color=CARD_WHITE, border_color=MINT_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
        
        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.15), width - Inches(0.36), Inches(0.4))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = title_color
        return card

    # =========================================================================
    # SLIDE 1: Title Slide (Mint & White)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_white_bg(s1)
    
    # Large Hero Card
    hero = add_white_card(s1, Inches(1.2), Inches(1.0), Inches(10.933), Inches(5.5), "", MINT_DARK, CARD_WHITE, MINT_BORDER)
    
    tb = s1.shapes.add_textbox(Inches(1.6), Inches(1.4), Inches(10.133), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = "CAPITOL UNIVERSITY • COLLEGE OF COMPUTER STUDIES"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = MINT_DARK
    p0.space_after = Pt(12)
    
    p1 = tf.add_paragraph()
    p1.text = "Warehouse Inventory Management System"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_DARK
    p1.space_after = Pt(6)
    
    p2 = tf.add_paragraph()
    p2.text = "Subsystem 1: Real-Time Stock Tracking, Quick Floor Audits & Valuation for Winzelle Store"
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEAL
    p2.space_after = Pt(24)
    
    p3 = tf.add_paragraph()
    p3.text = "IT Project 5 Final Defense Presentation"
    p3.font.size = Pt(13)
    p3.font.bold = True
    p3.font.color.rgb = TEXT_DARK
    p3.space_after = Pt(14)
    
    p4 = tf.add_paragraph()
    p4.text = (
        "Student / Presenter: Alexa A. Regodos\n"
        "Degree Program: Bachelor of Science in Information Technology\n"
        "Target Store: Winzelle Store (Wholesale Beverage & Goods Distribution)\n"
        "Date of Defense: October 2026\n"
        "System Version: Built with Laravel Framework 13.33.0 & React 19.2"
    )
    p4.font.size = Pt(12)
    p4.font.color.rgb = TEXT_MUTED

    s1.notes_slide.notes_text_frame.text = (
        "Good morning / afternoon members of the panel. Today I will present Subsystem 1: "
        "The Warehouse Inventory Management System for Winzelle Store. This system helps the store "
        "track real-time stocks, avoid running out of popular drinks, perform quick physical counts, "
        "and know the exact total money value of stocks currently on the shelves."
    )

    # =========================================================================
    # SLIDE 2: Agenda & Simple Overview
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_white_bg(s2)
    add_clean_header(s2, "Presentation Overview & Roadmap", "1. The Foundation")
    
    # Left Card: Summary in everyday language
    add_white_card(s2, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2), "What This System Does", MINT_DARK)
    tb = s2.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "Subsystem 1 is the central inventory system of Winzelle Store.\n\n"
        "• Replaces manual paper lists and messy notebooks with an easy-to-use web dashboard.\n\n"
        "• Shows real-time stocks for all products (like Pepsi, Coke, Sting, water, juices).\n\n"
        "• Lets staff update counts in one click right in the table, or change multiple items at once.\n\n"
        "• Automatically shows stock warnings: 'In Stock' (green), 'Low Stock' (orange), and 'Out of Stock' (red).\n\n"
        "• Automatically adds new stocks whenever a delivery arrives from distributors."
    )
    p.font.size = Pt(12.5)
    p.font.color.rgb = TEXT_MUTED

    # Right Card: 15-Slide Agenda
    add_white_card(s2, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2), "Presentation Roadmap", TEAL)
    tb = s2.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    steps = [
        "1. Store Background: How Winzelle Store operates today",
        "2. The Problem: Common stock errors and manual counting delays",
        "3. System Objectives: What the system achieves for the store",
        "4. How the System Works: Web browser, server, and database",
        "5. Database Tables: How products and quantities are stored",
        "6. Practical Rules: Calculating stock value and preventing negative stock",
        "7. Built With: Modern PHP 8.3, Laravel 13, and React 19",
        "8. Screen Walkthrough: Search bar, filter pills, batch buttons",
        "9. Actual Testing: Verifying that all features work properly",
        "10. Security & User Roles: Admin, Owner, and Checker permissions",
        "11. Conclusion & Next Step: Connecting to Subsystem 2 Deliveries"
    ]
    for i, step in enumerate(steps):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {step}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(3)

    # =========================================================================
    # SLIDE 3: Store Background & Daily Operations
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_white_bg(s3)
    add_clean_header(s3, "Store Background: Daily Operations at Winzelle", "1. The Foundation")
    
    c_w = Inches(3.65)
    c_gap = Inches(0.38)
    
    # Card 1: Fast-Moving Beverages
    add_white_card(s3, Inches(0.8), Inches(1.7), c_w, Inches(5.2), "Wholesale Beverages", TEAL)
    tb = s3.shapes.add_textbox(Inches(0.95), Inches(2.2), c_w - Inches(0.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "Winzelle Store sells wholesale beverages to sari-sari stores, small eateries, and walk-in buyers:\n\n"
        "• Products from major providers like Pepsi-Cola and Coca-Cola.\n\n"
        "• Various bottle types:\n"
        "  - Returnable Glass Bottles (e.g. Sting 240ml RGB/24)\n"
        "  - Plastic PET bottles (e.g. Pepsi 195ml PET/12, Coke 1.5L)\n"
        "  - Cans and water bottles.\n\n"
        "Cases move quickly every day, so stocks change every few hours."
    )
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED

    # Card 2: Physical Counts on the Floor
    add_white_card(s3, Inches(0.8) + c_w + c_gap, Inches(1.7), c_w, Inches(5.2), "Floor Audits & Counting", MINT_DARK)
    tb = s3.shapes.add_textbox(Inches(0.95) + c_w + c_gap, Inches(2.2), c_w - Inches(0.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "Warehouse checkers walk through the aisles to count remaining cases:\n\n"
        "• They check if physical cases on shelves match what is recorded.\n\n"
        "• They need to quickly note down broken bottles, spoilage, or sold items.\n\n"
        "• Previously, they wrote counts on paper pads, which took hours to type into Excel later."
    )
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED

    # Card 3: Stock Value & Re-ordering
    add_white_card(s3, Inches(0.8) + (c_w + c_gap) * 2, Inches(1.7), c_w, Inches(5.2), "Total Stock Value & Re-order", AMBER_TEXT)
    tb = s3.shapes.add_textbox(Inches(0.95) + (c_w + c_gap) * 2, Inches(2.2), c_w - Inches(0.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "The store owner needs to answer two key business questions every day:\n\n"
        "1. 'How much money is currently tied up in our warehouse inventory?'\n"
        "   (Total Valuation in ₱)\n\n"
        "2. 'Which items are about to run out so we can re-order from the distributor today?'\n\n"
        "Subsystem 1 displays both answers automatically."
    )
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 4: Problem Statement (Clear & Practical)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_white_bg(s4)
    add_clean_header(s4, "Problems Identified in the Existing Process", "1. The Foundation")
    
    g_w = Inches(5.6)
    g_h = Inches(2.45)
    
    # Box 1: Paper Delay
    add_white_card(s4, Inches(0.8), Inches(1.7), g_w, g_h, "1. Slow Paper Recording & Lost Notes", ROSE_TEXT, MINT_LIGHT, MINT_BORDER)
    tb = s4.shapes.add_textbox(Inches(1.0), Inches(2.15), g_w - Inches(0.4), g_h - Inches(0.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "When checkers write physical counts on paper, papers can get lost, wet, or misread. Transferring numbers to spreadsheets at the end of the day causes delays, leaving the owner with outdated stock numbers."
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_DARK

    # Box 2: Accidental Negative Numbers
    add_white_card(s4, Inches(6.9), Inches(1.7), g_w, g_h, "2. Human Errors & Negative Numbers", ROSE_TEXT, MINT_LIGHT, MINT_BORDER)
    tb = s4.shapes.add_textbox(Inches(7.1), Inches(2.15), g_w - Inches(0.4), g_h - Inches(0.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "In Excel, someone can easily type a wrong number or subtract too much, leading to impossible counts like '-5 cases'. This messes up the store's inventory value and financial calculations."
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_DARK

    # Box 3: Unexpected Out-of-Stock
    add_white_card(s4, Inches(0.8), Inches(4.45), g_w, g_h, "3. Running Out of Stock Without Warning", AMBER_TEXT, MINT_LIGHT, MINT_BORDER)
    tb = s4.shapes.add_textbox(Inches(1.0), Inches(4.9), g_w - Inches(0.4), g_h - Inches(0.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Without automated alerts, staff only notice that popular drinks like Coke 1.5L or Sting are empty when customers arrive to buy them. This results in lost sales and disappointed store customers."
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_DARK

    # Box 4: No History of Who Changed What
    add_white_card(s4, Inches(6.9), Inches(4.45), g_w, g_h, "4. No Record of Who Adjusted the Stock", TEAL, MINT_LIGHT, MINT_BORDER)
    tb = s4.shapes.add_textbox(Inches(7.1), Inches(4.9), g_w - Inches(0.4), g_h - Inches(0.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "When someone changes a number in a spreadsheet, the old number is gone forever. If there is a missing case of drinks, the owner cannot check who made the change, when it was changed, or why."
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 5: Objectives & Scope
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_white_bg(s5)
    add_clean_header(s5, "Objectives & What the System Can Do", "2. Research & Design")
    
    # Left: Objectives
    add_white_card(s5, Inches(0.8), Inches(1.7), Inches(6.5), Inches(5.2), "Main Goals of the Project", MINT_DARK)
    tb = s5.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(6.1), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "GENERAL OBJECTIVE:"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    p2 = tf.add_paragraph()
    p2.text = "To build a simple, reliable, web-based Warehouse Inventory Management System for Winzelle Store that tracks real-time stock balances, prevents counting errors, warns staff before items run out, and logs all changes."
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_after = Pt(10)
    
    p3 = tf.add_paragraph()
    p3.text = "SPECIFIC OBJECTIVES (WHAT IT DOES):"
    p3.font.size = Pt(12)
    p3.font.bold = True
    p3.font.color.rgb = TEXT_DARK
    
    objs = [
        "1. Central Stock Table: Show product name, SKU, distributor, case quantity, cost price, selling price, and total value.",
        "2. Live Search & Filters: Instantly search products and filter by brand (Pepsi, Coke) or category (Carbonated, Energy).",
        "3. Quick Stock Updating: Let staff edit a quantity in 1 click, or select 5 items and update all of them together.",
        "4. Automatic Low-Stock Warning: Highlight products in orange if 15 cases or fewer remain, and red if 0 cases.",
        "5. Excel Download: Download an official formatted Excel spreadsheet matching the store's green template.",
        "6. Inbound Delivery Link: Automatically increase stock whenever a delivery is saved in Subsystem 2."
    ]
    for o in objs:
        p = tf.add_paragraph()
        p.text = f"• {o}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(2)

    # Right: Scope & Boundaries
    add_white_card(s5, Inches(7.6), Inches(1.7), Inches(4.9), Inches(5.2), "Project Scope & Limitations", TEAL)
    tb = s5.shapes.add_textbox(Inches(7.8), Inches(2.2), Inches(4.5), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "WHAT IS INCLUDED:"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    in_items = [
        "Full inventory list for all beverage products.",
        "Inline edit and batch stock adjustments.",
        "Activity log that saves who updated each item.",
        "Excel (.xlsx) and CSV export.",
        "Role-based access: Admin, Owner, and Checker."
    ]
    for it in in_items:
        p = tf.add_paragraph()
        p.text = f"✓ {it}"
        p.font.size = Pt(11)
        p.font.color.rgb = MINT_DARK
        p.space_after = Pt(2)
        
    p = tf.add_paragraph()
    p.text = "\nWHAT IS NOT INCLUDED (FOR NOW):"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    out_items = [
        "Automatic warehouse conveyor belts or robotics.",
        "Mobile App Store apps (works directly in mobile browser).",
        "Multi-currency (everything is in Philippine Pesos ₱)."
    ]
    for it in out_items:
        p = tf.add_paragraph()
        p.text = f"✗ {it}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_SUBTLE
        p.space_after = Pt(2)

    # =========================================================================
    # SLIDE 6: How the System Works (Architecture)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_white_bg(s6)
    add_clean_header(s6, "How the System Works: Simple 3-Step Flow", "2. Research & Design")
    
    step_w = Inches(3.65)
    
    # Step 1: User Screen
    add_white_card(s6, Inches(0.8), Inches(1.7), step_w, Inches(5.2), "1. Staff Screen (Frontend)", MINT_DARK)
    tb = s6.shapes.add_textbox(Inches(0.95), Inches(2.2), step_w - Inches(0.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "WHAT THE USER SEES & USES:\n\n"
        "• Opened in Google Chrome or any browser on a PC, tablet, or phone.\n\n"
        "• Built with React 19 and Tailwind CSS.\n\n"
        "• Displays the clean table, search box, summary cards, and buttons.\n\n"
        "• When you click or type, the screen updates immediately without making you wait or reloading the whole page."
    )
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED

    # Step 2: The Engine (Backend)
    add_white_card(s6, Inches(0.8) + step_w + c_gap, Inches(1.7), step_w, Inches(5.2), "2. The System Engine (Server)", TEAL)
    tb = s6.shapes.add_textbox(Inches(0.95) + step_w + c_gap, Inches(2.2), step_w - Inches(0.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "HOW IT PROCESSES REQUESTS:\n\n"
        "• Powered by Laravel 13 running on PHP 8.3.\n\n"
        "• Checks who is logged in (Is it an Admin, Owner, or Checker?).\n\n"
        "• Enforces rules: Rejects negative numbers and validates stock edits.\n\n"
        "• Connects Subsystem 1 (Inventory) with Subsystem 2 (Deliveries)."
    )
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED

    # Step 3: The Safe Storage (Database)
    add_white_card(s6, Inches(0.8) + (step_w + c_gap) * 2, Inches(1.7), step_w, Inches(5.2), "3. Safe Storage (MySQL)", MINT_DARK)
    tb = s6.shapes.add_textbox(Inches(0.95) + (step_w + c_gap) * 2, Inches(2.2), step_w - Inches(0.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "WHERE DATA IS SAFELY STORED:\n\n"
        "• Stored in MySQL database named 'subsystem2'.\n\n"
        "• Keeps product quantities, prices, user accounts, and activity history.\n\n"
        "• Safe saving: Changes are permanently written to disk so numbers are never lost even if the PC restarts."
    )
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 7: Database Design (Table of Records)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_white_bg(s7)
    add_clean_header(s7, "Database Design: How Product Stocks Are Stored", "2. Research & Design")
    
    # Left: inventories Table
    add_white_card(s7, Inches(0.8), Inches(1.7), Inches(7.5), Inches(5.2), "The Primary Storage: 'inventories' Table", MINT_DARK)
    
    rows, cols = 8, 4
    table_shape = s7.shapes.add_table(rows, cols, Inches(1.0), Inches(2.3), Inches(7.1), Inches(4.2))
    table = table_shape.table
    table.columns[0].width = Inches(1.5)
    table.columns[1].width = Inches(1.4)
    table.columns[2].width = Inches(1.0)
    table.columns[3].width = Inches(3.2)
    
    t_heads = ["Field Name", "Type", "Required?", "What It Stores"]
    for c_idx, h in enumerate(t_heads):
        cell = table.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = MINT_LIGHT
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = MINT_DARK
        
    t_rows_data = [
        ("id", "Number", "Yes", "Unique record ID for each item"),
        ("product_id", "Number", "Yes", "Connects to product catalog"),
        ("sku", "Text", "Optional", "Short code like 'PEP-101'"),
        ("category", "Text", "Yes", "e.g. 'Carbonated Soft Drinks'"),
        ("product_name", "Text", "Yes", "e.g. 'Pepsi Reg 195ml PET/12'"),
        ("quantity", "Number", "Yes", "Current case count (cannot be negative)"),
        ("purchase_price", "Decimal (₱)", "Yes", "Supplier buying cost (e.g. ₱114.00)")
    ]
    for r_idx, r_data in enumerate(t_rows_data):
        for c_idx, val in enumerate(r_data):
            cell = table.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_WHITE
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10.5)
            p.font.color.rgb = TEXT_DARK

    # Right: How Tables Relate
    add_white_card(s7, Inches(8.6), Inches(1.7), Inches(3.9), Inches(5.2), "Connected Tables", TEAL)
    tb = s7.shapes.add_textbox(Inches(8.8), Inches(2.2), Inches(3.5), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    pts = [
        "Products Table: Stores master product list and selling prices.",
        "Purchases Table (Subsystem 2): Stores delivery receipts from Pepsi and Coke.",
        "Activity Logs Table: Saves a record every time stock is edited (who, when, old qty, new qty).",
        "Settings Table: Saves the store's low stock alert threshold (default: 15 cases)."
    ]
    for pt in pts:
        p = tf.add_paragraph()
        p.text = f"• {pt}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(8)

    # =========================================================================
    # SLIDE 8: System Logic & Formulas (Simple & Relatable)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_white_bg(s8)
    add_clean_header(s8, "System Logic: Formulas & Rules Made Simple", "2. Research & Design")
    
    # 3 Cards
    add_white_card(s8, Inches(0.8), Inches(1.7), step_w, Inches(5.2), "1. Stock Value Calculation", MINT_DARK)
    tb = s8.shapes.add_textbox(Inches(0.95), Inches(2.2), step_w - Inches(0.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "HOW WE COMPUTE STOCK VALUE:\n\n"
        "Value per Item =\n"
        "  Quantity × Purchase Price\n\n"
        "Total Warehouse Value =\n"
        "  Sum of all items\n\n"
        "REAL EXAMPLE:\n"
        "• 50 cases of Pepsi at ₱114.00\n"
        "  = ₱5,700.00\n"
        "• 100 cases of Sting at ₱280.00\n"
        "  = ₱28,000.00\n\n"
        "When exported to Excel, it uses real formulas (=E4*G4) so numbers recalculate automatically."
    )
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED

    add_white_card(s8, Inches(0.8) + step_w + c_gap, Inches(1.7), step_w, Inches(5.2), "2. Preventing Negative Stocks", ROSE_TEXT)
    tb = s8.shapes.add_textbox(Inches(0.95) + step_w + c_gap, Inches(2.2), step_w - Inches(0.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "THE ZERO-FLOOR RULE:\n\n"
        "New Stock = Maximum of (0, Current Stock ± Change)\n\n"
        "WHY THIS MATTERS:\n"
        "In physical reality, you can have 0 bottles, but you cannot have negative bottles on a shelf.\n\n"
        "EXAMPLE:\n"
        "If you have 5 cases left and accidentally batch-subtract 8 cases:\n"
        "• The system sets it to 0.\n"
        "• It never allows '-3'.\n"
        "• It immediately sends an 'Out of Stock' alert."
    )
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED

    add_white_card(s8, Inches(0.8) + (step_w + c_gap) * 2, Inches(1.7), step_w, Inches(5.2), "3. Automatic Alert Badges", AMBER_TEXT)
    tb = s8.shapes.add_textbox(Inches(0.95) + (step_w + c_gap) * 2, Inches(2.2), step_w - Inches(0.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "THREE COLOR-CODED STATUSES:\n\n"
        "• GREEN: 'In Stock'\n"
        "  More than 15 cases remaining.\n"
        "  Stock is healthy.\n\n"
        "• ORANGE: 'Low Stock'\n"
        "  15 cases or fewer remaining.\n"
        "  Alerts the owner to place a re-order soon.\n\n"
        "• RED: 'Out of Stock'\n"
        "  0 cases remaining.\n"
        "  Urgent alert to contact distributor."
    )
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 9: Technology Stack (Exact Versions from Workspace)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_white_bg(s9)
    add_clean_header(s9, "Implementation: Exact Tools Used in This Project", "3. Execution & Validation")
    
    # 2 Comparison Cards
    add_white_card(s9, Inches(0.8), Inches(1.7), g_w, Inches(5.2), "Backend Tools & Version Check", MINT_DARK)
    tb = s9.shapes.add_textbox(Inches(1.0), Inches(2.2), g_w - Inches(0.4), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    b_tools = [
        ("Laravel Framework 13.33.0", "The main PHP framework running our server and web pages (checked via 'php artisan --version')."),
        ("PHP 8.3.35", "Fast, modern PHP runtime installed on the XAMPP workstation."),
        ("Inertia.js v3.0.0", "Connects Laravel directly to React components without needing complex API code."),
        ("MySQL 8.0+ Database", "Stores our 'subsystem2' database with tables for products, inventory, and activity logs.")
    ]
    for title, desc in b_tools:
        p = tf.add_paragraph()
        p.text = f"• {title}:"
        p.font.bold = True
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_DARK
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_after = Pt(4)

    add_white_card(s9, Inches(6.9), Inches(1.7), g_w, Inches(5.2), "Frontend Tools & Version Check", TEAL)
    tb = s9.shapes.add_textbox(Inches(7.1), Inches(2.2), g_w - Inches(0.4), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    f_tools = [
        ("React 19.2.0", "Builds our interactive single-page user interface with instant button responses."),
        ("TypeScript 5.7.2", "Adds type checking to JavaScript to prevent coding typos and bugs."),
        ("Tailwind CSS 4.0.0", "Provides our clean mint-green and white responsive styling."),
        ("ExcelJS ^4.4.0", "Creates real formatted Microsoft Excel spreadsheets directly inside the browser."),
        ("Node.js v22.17.0 & Vite 8.0.0", "Compiles and loads our frontend in milliseconds.")
    ]
    for title, desc in f_tools:
        p = tf.add_paragraph()
        p.text = f"• {title}:"
        p.font.bold = True
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_DARK
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_after = Pt(4)

    # =========================================================================
    # SLIDE 10: Screen Demonstration & Features
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_white_bg(s10)
    add_clean_header(s10, "System Demonstration: Walkthrough of the Real UI", "3. Execution & Validation")
    
    # 3 Cards representing UI parts
    add_white_card(s10, Inches(0.8), Inches(1.7), step_w, Inches(5.2), "1. Top Summary Banner", MINT_DARK)
    tb = s10.shapes.add_textbox(Inches(0.95), Inches(2.2), step_w - Inches(0.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "THREE SUMMARY CARDS AT THE TOP:\n\n"
        "• Total Product SKUs:\n"
        "  Shows total items in the store.\n\n"
        "• Total Units in Stock:\n"
        "  Sums up all physical cases across all drinks.\n\n"
        "• Total Inventory Valuation:\n"
        "  Calculates the total monetary value in Philippine Pesos (₱) right now."
    )
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED

    add_white_card(s10, Inches(0.8) + step_w + c_gap, Inches(1.7), step_w, Inches(5.2), "2. Search Bar & Filter Pills", TEAL)
    tb = s10.shapes.add_textbox(Inches(0.95) + step_w + c_gap, Inches(2.2), step_w - Inches(0.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "FINDING PRODUCTS QUICKLY:\n\n"
        "• Live Search Bar:\n"
        "  Type 'Pepsi' or 'Coke' and the table filters instantly as you type.\n\n"
        "• Multi-Select Distributor Pills:\n"
        "  Click 'Pepsi' or 'Coca-Cola' pills to see only that brand.\n\n"
        "• Category Pills:\n"
        "  Filter by 'Carbonated Soft Drinks', 'Energy Drinks', or 'Water'."
    )
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED

    add_white_card(s10, Inches(0.8) + (step_w + c_gap) * 2, Inches(1.7), step_w, Inches(5.2), "3. Editing & Batch Actions", MINT_DARK)
    tb = s10.shapes.add_textbox(Inches(0.95) + (step_w + c_gap) * 2, Inches(2.2), step_w - Inches(0.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "UPDATING & EXPORTING STOCKS:\n\n"
        "• 1-Click Inline Edit:\n"
        "  Click any stock number on the table, type the new count, click checkmark.\n\n"
        "• Batch Toolbar:\n"
        "  Select 5 checkboxes -> Click 'Adjust Stock' -> Add +10 cases to all 5 at once.\n\n"
        "• Export Buttons:\n"
        "  One click to download formatted Excel or CSV report."
    )
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 11: System Testing (Simple Test Cases)
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_white_bg(s11)
    add_clean_header(s11, "System Testing: Verifying That Features Work", "3. Execution & Validation")
    
    add_white_card(s11, Inches(0.8), Inches(1.7), Inches(11.733), Inches(5.2), "Sample Test Cases Executed on the Real System", MINT_DARK)
    
    t_shape11 = s11.shapes.add_table(8, 5, Inches(1.0), Inches(2.3), Inches(11.333), Inches(4.2))
    table11 = t_shape11.table
    table11.columns[0].width = Inches(1.0)
    table11.columns[1].width = Inches(2.7)
    table11.columns[2].width = Inches(3.2)
    table11.columns[3].width = Inches(3.4)
    table11.columns[4].width = Inches(1.0)
    
    t11_h = ["Test No.", "Action Tested", "What Was Done", "Expected & Actual Result", "Result"]
    for c_idx, h in enumerate(t11_h):
        cell = table11.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = MINT_LIGHT
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = MINT_DARK
        
    tests_data = [
        ("Test 1", "Load Inventory Page", "Open /inventory in browser", "Shows all beverage products with prices and stock counts", "PASS"),
        ("Test 2", "Live Search Bar", "Type 'Pepsi' in search box", "Table instantly narrows down to Pepsi items within milliseconds", "PASS"),
        ("Test 3", "1-Click Inline Edit", "Change stock from 10 to 20", "Stock updates immediately; notification dispatched to all users", "PASS"),
        ("Test 4", "Prevent Negative Stock", "Try entering -5 in quantity", "System rejects the negative input and keeps the original number", "PASS"),
        ("Test 5", "Batch Subtract Delta", "Subtract 50 cases from 15 stock", "Quantity safely stops at 0; 'Out of Stock' alert is sent", "PASS"),
        ("Test 6", "Excel Spreadsheet Export", "Click 'Export Excel' button", "Downloads a green-styled .xlsx file with working total formulas", "PASS"),
        ("Test 7", "Delivery Inbound Sync", "Add 25 cases in Purchase page", "Inventory table automatically increments by +25 cases", "PASS")
    ]
    for r_idx, row in enumerate(tests_data):
        for c_idx, val in enumerate(row):
            cell = table11.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_WHITE
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10.5)
            if c_idx == 4:
                p.font.bold = True
                p.font.color.rgb = MINT_DARK
            else:
                p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 12: Security, User Roles & Activity Log
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_white_bg(s12)
    add_clean_header(s12, "Security & Accountability: User Roles and Logs", "3. Execution & Validation")
    
    # 3 Cards
    add_white_card(s12, Inches(0.8), Inches(1.7), step_w, Inches(5.2), "1. User Roles & Permissions", MINT_DARK)
    tb = s12.shapes.add_textbox(Inches(0.95), Inches(2.2), step_w - Inches(0.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "THREE CLEAR ROLES:\n\n"
        "• Admin:\n"
        "  Full control: Can edit all stocks, run batch actions, view logs, and change settings.\n\n"
        "• Owner:\n"
        "  Can monitor stock value, adjust stocks, view activity logs, and export reports.\n\n"
        "• Checker:\n"
        "  Can view inventory and update counts during floor audits, but cannot change system settings."
    )
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED

    add_white_card(s12, Inches(0.8) + step_w + c_gap, Inches(1.7), step_w, Inches(5.2), "2. The Activity Audit Log", TEAL)
    tb = s12.shapes.add_textbox(Inches(0.95) + step_w + c_gap, Inches(2.2), step_w - Inches(0.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "TRACKING EVERY SINGLE CHANGE:\n\n"
        "Whenever stock is changed, the system saves a permanent record in 'activity_logs':\n\n"
        "• Who changed it (e.g. 'Checker John Doe')\n"
        "• What product was changed (e.g. 'Coke 1.5L')\n"
        "• What the old count was (e.g. 10 cases)\n"
        "• What the new count is (e.g. 25 cases)\n"
        "• Exact date and timestamp.\n\n"
        "Eliminates mystery stock loss."
    )
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED

    add_white_card(s12, Inches(0.8) + (step_w + c_gap) * 2, Inches(1.7), step_w, Inches(5.2), "3. Safe & Reliable Saving", MINT_DARK)
    tb = s12.shapes.add_textbox(Inches(0.95) + (step_w + c_gap) * 2, Inches(2.2), step_w - Inches(0.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "PROTECTING STORE DATA:\n\n"
        "• Transaction Safety:\n"
        "  When updating 10 items at once, either all 10 are saved successfully, or none are changed if an error occurs. No half-saved data.\n\n"
        "• Archive Safety (Soft Deletes):\n"
        "  Accidentally archived records are not destroyed; they can be restored in 1 click."
    )
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 13: Conclusion & Practical Benefits
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_white_bg(s13)
    add_clean_header(s13, "Conclusion: Real Benefits for Winzelle Store", "4. Defense & Conclusion")
    
    add_white_card(s13, Inches(0.8), Inches(1.7), Inches(11.733), Inches(5.2), "How Subsystem 1 Solves Daily Store Problems", MINT_DARK)
    tb = s13.shapes.add_textbox(Inches(1.1), Inches(2.2), Inches(11.1), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    benefits = [
        ("No More Paper Notes", "Checkers can bring a tablet or phone right down the warehouse aisle and enter physical counts in seconds."),
        ("Instant Stock Value (Valuation in ₱)", "The owner can check their screen at any time to see the total peso value of all stocks currently sitting on the shelves."),
        ("Early Warning Before Drinks Run Out", "The orange 'Low Stock' badge gives staff advance notice to place orders with Pepsi or Coke before shelves are empty."),
        ("Full Accountability with Audit Logs", "Every single stock addition or reduction has the staff name and timestamp attached to it."),
        ("Clean Excel Reports for Tax & Bookkeeping", "Generates professional green-styled Excel spreadsheets with formulas in one click."),
        ("Automated Delivery Intake", "Incoming delivery receipts in Subsystem 2 automatically add stock to Subsystem 1, saving time and avoiding double entry.")
    ]
    for i, (b_title, b_desc) in enumerate(benefits):
        p = tf.add_paragraph()
        p.text = f"{i+1}. {b_title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        
        p2 = tf.add_paragraph()
        p2.text = b_desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_after = Pt(4)

    # =========================================================================
    # SLIDE 14: Limitations & Future Plans
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    set_white_bg(s14)
    add_clean_header(s14, "Current Limitations & Future Recommendations", "4. Defense & Conclusion")
    
    add_white_card(s14, Inches(0.8), Inches(1.7), g_w, Inches(5.2), "Current System Limitations", AMBER_TEXT)
    tb = s14.shapes.add_textbox(Inches(1.0), Inches(2.2), g_w - Inches(0.4), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    lims = [
        ("Typing vs Barcode Scanning", "Staff currently type counts using a keyboard, tablet screen, or handheld barcode scanner acting as a keyboard."),
        ("Single Currency", "All calculations and valuation numbers are in Philippine Pesos (₱)."),
        ("Single Store Location", "Currently built for one main warehouse/store; cross-branch transfers between multiple buildings can be added later.")
    ]
    for title, desc in lims:
        p = tf.add_paragraph()
        p.text = f"• {title}:"
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_after = Pt(6)

    add_white_card(s14, Inches(6.9), Inches(1.7), g_w, Inches(5.2), "Future Development Plans", TEAL)
    tb = s14.shapes.add_textbox(Inches(7.1), Inches(2.2), g_w - Inches(0.4), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    plans = [
        ("Camera Barcode & QR Scanning", "Allow staff to scan bottle and crate barcodes directly using their phone's camera."),
        ("Automatic Re-order Reminders", "Suggest how many cases to order based on how fast Coke or Pepsi sold over the last 7 days."),
        ("Full Subsystem 2 Presentation", "Presenting the Distributor Onboarding, Pricing Margins, and BIR 2023 12% VAT calculations in Subsystem 2.")
    ]
    for title, desc in plans:
        p = tf.add_paragraph()
        p.text = f"✓ {title}:"
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_after = Pt(6)

    # =========================================================================
    # SLIDE 15: References & Thank You / Q&A
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    set_white_bg(s15)
    add_clean_header(s15, "References & Open for Questions", "4. Defense & Conclusion")
    
    # Left: References
    add_white_card(s15, Inches(0.8), Inches(1.7), Inches(6.2), Inches(5.2), "Project References", TEAL)
    tb = s15.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.8), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    ref_list = [
        "1. Laravel 13 Documentation: Database Transactions & Eloquent Models (2026).",
        "2. React 19 Official Documentation: Component State & Fast Re-rendering (2026).",
        "3. Inertia.js Official Guide: Connecting Server-Side Controllers with Single Page Apps (2026).",
        "4. Bureau of Internal Revenue (BIR) Philippines: VAT Calculation Guidelines for Wholesale Trade (2023).",
        "5. Winzelle Store Operational Records & Official Inventory Template (docs/template-inventory.xlsx)."
    ]
    for r in ref_list:
        p = tf.add_paragraph()
        p.text = r
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(8)

    # Right: Thank you
    add_white_card(s15, Inches(7.3), Inches(1.7), Inches(5.2), Inches(5.2), "Defense Evaluation", MINT_DARK)
    tb = s15.shapes.add_textbox(Inches(7.5), Inches(2.5), Inches(4.8), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Thank You!"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = MINT_DARK
    p.space_after = Pt(8)
    
    p = tf.add_paragraph()
    p.text = "The presentation for Subsystem 1 is concluded.\nI am now ready for questions and recommendations from the panel."
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(20)
    
    p = tf.add_paragraph()
    p.text = (
        "Presenter: Alexa A. Regodos\n"
        "College of Computer Studies — Capitol University\n"
        "Project: Subsystem 1 (Warehouse Inventory Management)"
    )
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED

    # Save files
    targets = [
        "docs/Final Presentation for IT Project 5.pptx",
        "docs/Final Presentation for IT Project 5 - Subsystem 1.pptx",
        "docs/Winzelle_Subsystem1_Mint_Presentation.pptx"
    ]
    saved = []
    for t in targets:
        try:
            prs.save(t)
            saved.append(t)
        except PermissionError:
            print(f"Notice: {t} is currently open in PowerPoint. Skipping locked file.")
    print("Successfully saved files:", saved)

if __name__ == "__main__":
    build_mint_presentation()
